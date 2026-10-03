"""DOCX extractor."""
from pathlib import Path

from docx import Document

from .base import BaseExtractor, ExtractedDocument


class DocxExtractor(BaseExtractor):
    """Handles .docx files."""

    def extract(self, file_path: Path) -> ExtractedDocument:
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        doc = Document(str(file_path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            rows = [[c.text.strip() for c in row.cells] for row in table.rows]
            if rows:
                header = "| " + " | ".join(rows[0]) + " |"
                sep = "| " + " | ".join(["---"] * len(rows[0])) + " |"
                body = "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
                parts.append(f"{header}\n{sep}\n{body}")

        content = "\n\n".join(parts)
        return ExtractedDocument(
            content=content,
            source_path=str(file_path),
            file_type=".docx",
            metadata={"char_count": len(content)},
        )